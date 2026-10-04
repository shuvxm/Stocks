import os
import warnings
from datetime import datetime

import numpy as np
import pandas as pd
import yfinance as yf
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

warnings.filterwarnings("ignore")


# ============================================================
# CONFIGURATION
# ============================================================

ETF_LIST = [
    "NIFTYBEES.NS",
    "BANKBEES.NS",
    "ITBEES.NS",
    "JUNIORBEES.NS",
    "MID150BEES.NS",
    "PHARMABEES.NS",
    "AUTOBEES.NS",
    "PSUBNKBEES.NS",
    "CPSEETF.NS",
    "CONSUMBEES.NS",
    "FMCGIETF.NS",
    "HEALTHY.NS",
    "MON100.NS",
    "GOLDBEES.NS",
    "SILVERBEES.NS",
]

PERIOD = "2y"
INTERVAL = "1d"

# Minimum liquidity filter
MIN_AVG_VOLUME = 10000

# Swing filters
MIN_ADX = 20
MIN_RSI = 50
MAX_RSI = 70
MIN_RELATIVE_VOLUME = 0.80

# Score threshold
QUALIFY_SCORE = 65

# Output
OUTPUT_DIR = "ETF_REPORTS"
os.makedirs(OUTPUT_DIR, exist_ok=True)


# ============================================================
# DATA DOWNLOAD
# ============================================================

def download_data(ticker):
    try:
        df = yf.download(
            ticker,
            period=PERIOD,
            interval=INTERVAL,
            auto_adjust=True,
            progress=False,
            threads=False,
        )

        if df is None or df.empty:
            return None

        # Handle yfinance MultiIndex
        if isinstance(df.columns, pd.MultiIndex):
            df.columns = df.columns.get_level_values(0)

        required = ["Open", "High", "Low", "Close", "Volume"]

        for col in required:
            if col not in df.columns:
                return None

        df = df[required].copy()

        for col in required:
            df[col] = pd.to_numeric(df[col], errors="coerce")

        df.dropna(subset=["Close"], inplace=True)

        return df

    except Exception as e:
        print(f"ERROR downloading {ticker}: {e}")
        return None


# ============================================================
# TECHNICAL INDICATORS
# ============================================================

def calculate_indicators(df):

    df = df.copy()

    close = df["Close"]
    high = df["High"]
    low = df["Low"]
    volume = df["Volume"]

    # --------------------------------------------------------
    # RETURNS
    # --------------------------------------------------------

    df["Return_1D"] = close.pct_change(1) * 100
    df["Return_5D"] = close.pct_change(5) * 100
    df["Return_1M"] = close.pct_change(21) * 100
    df["Return_3M"] = close.pct_change(63) * 100
    df["Return_6M"] = close.pct_change(126) * 100
    df["Return_12M"] = close.pct_change(252) * 100

    # --------------------------------------------------------
    # MOVING AVERAGES
    # --------------------------------------------------------

    df["EMA_9"] = close.ewm(span=9, adjust=False).mean()
    df["EMA_20"] = close.ewm(span=20, adjust=False).mean()
    df["EMA_21"] = close.ewm(span=21, adjust=False).mean()
    df["EMA_50"] = close.ewm(span=50, adjust=False).mean()

    df["SMA_20"] = close.rolling(20).mean()
    df["SMA_50"] = close.rolling(50).mean()
    df["SMA_100"] = close.rolling(100).mean()
    df["SMA_200"] = close.rolling(200).mean()

    # --------------------------------------------------------
    # RSI 14
    # --------------------------------------------------------

    delta = close.diff()

    gain = delta.clip(lower=0)
    loss = -delta.clip(upper=0)

    avg_gain = gain.ewm(
        alpha=1 / 14,
        min_periods=14,
        adjust=False
    ).mean()

    avg_loss = loss.ewm(
        alpha=1 / 14,
        min_periods=14,
        adjust=False
    ).mean()

    rs = avg_gain / avg_loss.replace(0, np.nan)

    df["RSI_14"] = 100 - (100 / (1 + rs))

    # --------------------------------------------------------
    # MACD
    # --------------------------------------------------------

    ema12 = close.ewm(span=12, adjust=False).mean()
    ema26 = close.ewm(span=26, adjust=False).mean()

    df["MACD"] = ema12 - ema26
    df["MACD_Signal"] = df["MACD"].ewm(
        span=9,
        adjust=False
    ).mean()

    df["MACD_Histogram"] = (
        df["MACD"] - df["MACD_Signal"]
    )

    # --------------------------------------------------------
    # TRUE RANGE / ATR
    # --------------------------------------------------------

    previous_close = close.shift(1)

    tr1 = high - low
    tr2 = (high - previous_close).abs()
    tr3 = (low - previous_close).abs()

    true_range = pd.concat(
        [tr1, tr2, tr3],
        axis=1
    ).max(axis=1)

    df["TR"] = true_range

    df["ATR_14"] = true_range.ewm(
        alpha=1 / 14,
        adjust=False,
        min_periods=14
    ).mean()

    df["ATR_Percent"] = (
        df["ATR_14"] / close
    ) * 100

    # --------------------------------------------------------
    # ADX / DI
    # --------------------------------------------------------

    up_move = high.diff()
    down_move = -low.diff()

    plus_dm = np.where(
        (up_move > down_move) & (up_move > 0),
        up_move,
        0
    )

    minus_dm = np.where(
        (down_move > up_move) & (down_move > 0),
        down_move,
        0
    )

    atr = df["ATR_14"]

    plus_di = (
        100
        * pd.Series(plus_dm, index=df.index)
        .ewm(alpha=1 / 14, adjust=False)
        .mean()
        / atr
    )

    minus_di = (
        100
        * pd.Series(minus_dm, index=df.index)
        .ewm(alpha=1 / 14, adjust=False)
        .mean()
        / atr
    )

    dx = (
        100
        * (plus_di - minus_di).abs()
        / (plus_di + minus_di).replace(0, np.nan)
    )

    df["PLUS_DI"] = plus_di
    df["MINUS_DI"] = minus_di

    df["ADX_14"] = dx.ewm(
        alpha=1 / 14,
        adjust=False
    ).mean()

    # --------------------------------------------------------
    # BOLLINGER BANDS
    # --------------------------------------------------------

    bb_middle = close.rolling(20).mean()
    bb_std = close.rolling(20).std()

    df["BB_Middle"] = bb_middle
    df["BB_Upper"] = bb_middle + (2 * bb_std)
    df["BB_Lower"] = bb_middle - (2 * bb_std)

    band_range = (
        df["BB_Upper"] - df["BB_Lower"]
    )

    df["BB_Width"] = (
        band_range / bb_middle
    ) * 100

    df["BB_Percent_B"] = (
        (close - df["BB_Lower"])
        / band_range.replace(0, np.nan)
    ) * 100

    # --------------------------------------------------------
    # STOCHASTIC
    # --------------------------------------------------------

    lowest_14 = low.rolling(14).min()
    highest_14 = high.rolling(14).max()

    stoch_range = (
        highest_14 - lowest_14
    ).replace(0, np.nan)

    df["Stochastic_K"] = (
        (close - lowest_14)
        / stoch_range
    ) * 100

    df["Stochastic_D"] = (
        df["Stochastic_K"]
        .rolling(3)
        .mean()
    )

    # --------------------------------------------------------
    # ROC
    # --------------------------------------------------------

    df["ROC_14"] = (
        close.pct_change(14)
    ) * 100

    # --------------------------------------------------------
    # VOLUME
    # --------------------------------------------------------

    df["Volume_SMA_20"] = volume.rolling(20).mean()

    df["Relative_Volume"] = (
        volume
        / df["Volume_SMA_20"].replace(0, np.nan)
    )

    # --------------------------------------------------------
    # OBV
    # --------------------------------------------------------

    direction = np.sign(close.diff()).fillna(0)

    df["OBV"] = (
        direction * volume
    ).cumsum()

    df["OBV_EMA_20"] = (
        df["OBV"]
        .ewm(span=20, adjust=False)
        .mean()
    )

    # --------------------------------------------------------
    # PRICE LEVELS
    # --------------------------------------------------------

    df["High_20D"] = high.rolling(20).max()
    df["Low_20D"] = low.rolling(20).min()

    df["High_50D"] = high.rolling(50).max()
    df["Low_50D"] = low.rolling(50).min()

    df["High_52W"] = high.rolling(252).max()
    df["Low_52W"] = low.rolling(252).min()

    df["Distance_52W_High"] = (
        (close / df["High_52W"]) - 1
    ) * 100

    df["Distance_52W_Low"] = (
        (close / df["Low_52W"]) - 1
    ) * 100

    # --------------------------------------------------------
    # VOLATILITY
    # --------------------------------------------------------

    df["Volatility_20D"] = (
        close.pct_change()
        .rolling(20)
        .std()
        * np.sqrt(252)
        * 100
    )

    # --------------------------------------------------------
    # DRAWDOWN
    # --------------------------------------------------------

    rolling_max = close.cummax()

    df["Drawdown"] = (
        (close / rolling_max) - 1
    ) * 100

    return df


# ============================================================
# SCORE
# ============================================================

def calculate_score(row):

    score = 0

    # --------------------------------------------------------
    # TREND - 25 points
    # --------------------------------------------------------

    if row["Close"] > row["EMA_20"]:
        score += 5

    if row["EMA_20"] > row["EMA_50"]:
        score += 5

    if row["EMA_50"] > row["SMA_200"]:
        score += 5

    if row["Close"] > row["SMA_200"]:
        score += 5

    if row["SMA_50"] > row["SMA_200"]:
        score += 5

    # --------------------------------------------------------
    # MOMENTUM - 25 points
    # --------------------------------------------------------

    rsi = row["RSI_14"]

    if 50 <= rsi <= 70:
        score += 10
    elif 45 <= rsi < 50:
        score += 5

    if row["MACD"] > row["MACD_Signal"]:
        score += 10

    if row["ROC_14"] > 0:
        score += 5

    # --------------------------------------------------------
    # TREND STRENGTH - 15 points
    # --------------------------------------------------------

    if row["ADX_14"] >= 25:
        score += 10
    elif row["ADX_14"] >= 20:
        score += 5

    if row["PLUS_DI"] > row["MINUS_DI"]:
        score += 5

    # --------------------------------------------------------
    # VOLUME - 10 points
    # --------------------------------------------------------

    if row["Relative_Volume"] >= 1:
        score += 5

    if row["OBV"] > row["OBV_EMA_20"]:
        score += 5

    # --------------------------------------------------------
    # BREAKOUT - 10 points
    # --------------------------------------------------------

    if row["Close"] >= row["High_20D"] * 0.98:
        score += 5

    if row["Close"] >= row["High_50D"] * 0.95:
        score += 5

    # --------------------------------------------------------
    # RETURNS - 10 points
    # --------------------------------------------------------

    if row["Return_3M"] > 0:
        score += 5

    if row["Return_6M"] > 0:
        score += 5

    # --------------------------------------------------------
    # RISK PENALTY
    # --------------------------------------------------------

    if row["ATR_Percent"] > 5:
        score -= 5

    if row["Drawdown"] < -20:
        score -= 5

    return max(0, min(100, score))


# ============================================================
# CLASSIFICATION
# ============================================================

def classify(row):

    score = row["Score"]

    if score >= 80:
        return "STRONG SETUP"

    if score >= 65:
        return "QUALIFIED"

    if score >= 50:
        return "WATCH"

    return "FILTERED"


def trend_status(row):

    if (
        row["Close"] > row["EMA_20"]
        and row["EMA_20"] > row["EMA_50"]
        and row["EMA_50"] > row["SMA_200"]
    ):
        return "BULLISH"

    if (
        row["Close"] < row["EMA_20"]
        and row["EMA_20"] < row["EMA_50"]
        and row["EMA_50"] < row["SMA_200"]
    ):
        return "BEARISH"

    return "NEUTRAL"


def momentum_status(row):

    if (
        row["RSI_14"] >= 50
        and row["MACD"] > row["MACD_Signal"]
    ):
        return "POSITIVE"

    if (
        row["RSI_14"] < 45
        and row["MACD"] < row["MACD_Signal"]
    ):
        return "NEGATIVE"

    return "NEUTRAL"


def risk_status(row):

    atr = row["ATR_Percent"]
    drawdown = row["Drawdown"]

    if atr <= 2 and drawdown >= -10:
        return "LOW"

    if atr <= 4 and drawdown >= -20:
        return "MEDIUM"

    return "HIGH"


# ============================================================
# ANALYZE ONE ETF
# ============================================================

def analyze_etf(ticker):

    print(f"Analyzing {ticker}...")

    df = download_data(ticker)

    if df is None:
        return None, None

    if len(df) < 220:
        print(f"Not enough data for {ticker}")
        return None, None

    df = calculate_indicators(df)

    row = df.iloc[-1].copy()

    # --------------------------------------------------------
    # SCORE
    # --------------------------------------------------------

    row["Score"] = calculate_score(row)

    row["Signal"] = classify(row)
    row["Trend"] = trend_status(row)
    row["Momentum"] = momentum_status(row)
    row["Risk"] = risk_status(row)

    # --------------------------------------------------------
    # LIQUIDITY
    # --------------------------------------------------------

    row["Avg_Volume_20D"] = (
        df["Volume"]
        .rolling(20)
        .mean()
        .iloc[-1]
    )

    if row["Avg_Volume_20D"] < MIN_AVG_VOLUME:
        row["Liquidity_Filter"] = "FAIL"
    else:
        row["Liquidity_Filter"] = "PASS"

    # --------------------------------------------------------
    # SWING FILTERS
    # --------------------------------------------------------

    conditions = {
        "Price > EMA20":
            row["Close"] > row["EMA_20"],

        "EMA20 > EMA50":
            row["EMA_20"] > row["EMA_50"],

        "Price > SMA200":
            row["Close"] > row["SMA_200"],

        "RSI 50-70":
            MIN_RSI <= row["RSI_14"] <= MAX_RSI,

        "MACD Bullish":
            row["MACD"] > row["MACD_Signal"],

        "ADX >= 20":
            row["ADX_14"] >= MIN_ADX,

        "DI Bullish":
            row["PLUS_DI"] > row["MINUS_DI"],

        "Relative Volume":
            row["Relative_Volume"] >= MIN_RELATIVE_VOLUME,

        "3M Return Positive":
            row["Return_3M"] > 0,

        "Liquidity":
            row["Avg_Volume_20D"] >= MIN_AVG_VOLUME,
    }

    passed = sum(conditions.values())
    total = len(conditions)

    row["Filter_Passed"] = passed
    row["Filter_Total"] = total
    row["Filter_Percent"] = (
        passed / total
    ) * 100

    row["Qualified"] = (
        row["Score"] >= QUALIFY_SCORE
        and passed >= 7
    )

    # --------------------------------------------------------
    # TICKER
    # --------------------------------------------------------

    row["Ticker"] = ticker.replace(".NS", "")

    # --------------------------------------------------------
    # DATE
    # --------------------------------------------------------

    row["Analysis_Date"] = df.index[-1]

    return row, df


# ============================================================
# MAIN
# ============================================================

def main():

    print("=" * 70)
    print("INDIAN ETF SWING ANALYST")
    print("=" * 70)

    results = []
    historical_data = {}
    failed = []

    for ticker in ETF_LIST:

        try:

            result, history = analyze_etf(ticker)

            if result is not None:

                results.append(result)
                historical_data[ticker] = history

        except Exception as e:

            print(f"FAILED {ticker}: {e}")
            failed.append(ticker)

    if not results:

        print("No ETF data available.")
        return

    # --------------------------------------------------------
    # RESULTS DATAFRAME
    # --------------------------------------------------------

    results_df = pd.DataFrame(results)

    # Sort by score
    results_df.sort_values(
        by=["Qualified", "Score"],
        ascending=[False, False],
        inplace=True
    )

    # --------------------------------------------------------
    # SELECT COLUMNS
    # --------------------------------------------------------

    output_columns = [

        "Ticker",
        "Analysis_Date",
        "Close",

        "Score",
        "Signal",
        "Qualified",

        "Trend",
        "Momentum",
        "Risk",

        "RSI_14",

        "MACD",
        "MACD_Signal",
        "MACD_Histogram",

        "ADX_14",
        "PLUS_DI",
        "MINUS_DI",

        "EMA_9",
        "EMA_20",
        "EMA_21",
        "EMA_50",

        "SMA_20",
        "SMA_50",
        "SMA_100",
        "SMA_200",

        "ATR_14",
        "ATR_Percent",

        "BB_Upper",
        "BB_Middle",
        "BB_Lower",
        "BB_Percent_B",
        "BB_Width",

        "Stochastic_K",
        "Stochastic_D",

        "ROC_14",

        "Volume",
        "Volume_SMA_20",
        "Relative_Volume",

        "OBV",
        "OBV_EMA_20",

        "High_20D",
        "Low_20D",

        "High_50D",
        "Low_50D",

        "High_52W",
        "Low_52W",

        "Distance_52W_High",
        "Distance_52W_Low",

        "Return_1D",
        "Return_5D",
        "Return_1M",
        "Return_3M",
        "Return_6M",
        "Return_12M",

        "Volatility_20D",
        "Drawdown",

        "Avg_Volume_20D",

        "Filter_Passed",
        "Filter_Total",
        "Filter_Percent",

        "Liquidity_Filter",
    ]

    available_columns = [
        c for c in output_columns
        if c in results_df.columns
    ]

    all_results = results_df[available_columns].copy()

    # --------------------------------------------------------
    # QUALIFIED ETF TABLE
    # --------------------------------------------------------

    qualified = results_df[
        results_df["Qualified"] == True
    ].copy()

    qualified.sort_values(
        by="Score",
        ascending=False,
        inplace=True
    )

    qualified_columns = [
        "Ticker",
        "Analysis_Date",
        "Close",
        "Score",
        "Signal",
        "Trend",
        "Momentum",
        "Risk",
        "RSI_14",
        "MACD",
        "ADX_14",
        "ATR_Percent",
        "Relative_Volume",
        "Return_1M",
        "Return_3M",
        "Return_6M",
        "Distance_52W_High",
        "Filter_Passed",
        "Filter_Total",
    ]

    qualified_columns = [
        c for c in qualified_columns
        if c in qualified.columns
    ]

    qualified_df = qualified[
        qualified_columns
    ].copy()

    # --------------------------------------------------------
    # FILTERED OUT
    # --------------------------------------------------------

    filtered_df = results_df[
        results_df["Qualified"] == False
    ].copy()

    filtered_df.sort_values(
        by="Score",
        ascending=False,
        inplace=True
    )

    filtered_columns = [
        "Ticker",
        "Close",
        "Score",
        "Signal",
        "Trend",
        "Momentum",
        "Risk",
        "RSI_14",
        "MACD",
        "ADX_14",
        "Relative_Volume",
        "Return_3M",
        "Filter_Passed",
        "Filter_Total",
        "Filter_Percent",
    ]

    filtered_columns = [
        c for c in filtered_columns
        if c in filtered_df.columns
    ]

    filtered_df = filtered_df[
        filtered_columns
    ]

    # --------------------------------------------------------
    # CREATE EXCEL FILE
    # --------------------------------------------------------

    today = datetime.now().strftime("%Y-%m-%d")

    filename = os.path.join(
        OUTPUT_DIR,
        f"ETF_Swing_Analysis_{today}.xlsx"
    )

    with pd.ExcelWriter(
        filename,
        engine="openpyxl"
    ) as writer:

        qualified_df.to_excel(
            writer,
            sheet_name="Qualified ETFs",
            index=False
        )

        all_results.to_excel(
            writer,
            sheet_name="All ETFs",
            index=False
        )

        filtered_df.to_excel(
            writer,
            sheet_name="Filtered ETFs",
            index=False
        )

        # ----------------------------------------------------
        # DAILY PRICE DATA
        # ----------------------------------------------------

        price_rows = []

        for ticker, df in historical_data.items():

            temp = df.copy()
            temp["Ticker"] = ticker.replace(".NS", "")
            temp.reset_index(inplace=True)

            price_rows.append(temp)

        if price_rows:

            price_df = pd.concat(
                price_rows,
                ignore_index=True
            )

            price_df.to_excel(
                writer,
                sheet_name="Price History",
                index=False
            )

        # ----------------------------------------------------
        # CONFIGURATION
        # ----------------------------------------------------

        config = pd.DataFrame({
            "Setting": [
                "Period",
                "Interval",
                "Minimum ADX",
                "Minimum RSI",
                "Maximum RSI",
                "Minimum Relative Volume",
                "Minimum Average Volume",
                "Qualification Score",
            ],

            "Value": [
                PERIOD,
                INTERVAL,
                MIN_ADX,
                MIN_RSI,
                MAX_RSI,
                MIN_RELATIVE_VOLUME,
                MIN_AVG_VOLUME,
                QUALIFY_SCORE,
            ]
        })

        config.to_excel(
            writer,
            sheet_name="Configuration",
            index=False
        )

        # ----------------------------------------------------
        # FAILED TICKERS
        # ----------------------------------------------------

        pd.DataFrame({
            "Failed_Ticker": failed
        }).to_excel(
            writer,
            sheet_name="Failed",
            index=False
        )

    # --------------------------------------------------------
    # FORMAT EXCEL
    # --------------------------------------------------------

    format_excel(filename)

    # --------------------------------------------------------
    # CONSOLE OUTPUT
    # --------------------------------------------------------

    print()
    print("=" * 70)
    print("RESULT")
    print("=" * 70)

    print(f"Total ETFs analyzed : {len(results_df)}")
    print(f"Qualified ETFs      : {len(qualified_df)}")
    print(f"Filtered ETFs       : {len(filtered_df)}")
    print()

    if not qualified_df.empty:

        print("QUALIFIED ETF SETUPS")
        print("-" * 70)

        for _, r in qualified_df.iterrows():

            print(
                f"{r['Ticker']:15}"
                f" Price={r['Close']:.2f} "
                f"Score={r['Score']:.0f} "
                f"RSI={r['RSI_14']:.1f} "
                f"ADX={r['ADX_14']:.1f} "
                f"RVOL={r['Relative_Volume']:.2f} "
                f"{r['Signal']}"
            )

    else:

        print("No ETFs passed the qualification filter today.")

    print()
    print(f"Excel report: {filename}")
    print("=" * 70)


# ============================================================
# EXCEL FORMATTING
# ============================================================

def format_excel(filename):

    wb = load_workbook(filename)

    header_fill = PatternFill(
        "solid",
        fgColor="1F4E78"
    )

    header_font = Font(
        color="FFFFFF",
        bold=True
    )

    qualified_fill = PatternFill(
        "solid",
        fgColor="E2F0D9"
    )

    filtered_fill = PatternFill(
        "solid",
        fgColor="FCE4D6"
    )

    for ws in wb.worksheets:

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        # Header
        for cell in ws[1]:

            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center"
            )

        # Width
        for column_cells in ws.columns:

            max_length = 0

            for cell in column_cells:

                try:

                    value_length = len(
                        str(cell.value)
                    )

                    max_length = max(
                        max_length,
                        value_length
                    )

                except Exception:
                    pass

            width = min(
                max(max_length + 2, 10),
                28
            )

            column_letter = get_column_letter(
                column_cells[0].column
            )

            ws.column_dimensions[
                column_letter
            ].width = width

        # Number formatting
        for row in ws.iter_rows(
            min_row=2
        ):

            for cell in row:

                if isinstance(
                    cell.value,
                    (float, int)
                ):

                    cell.number_format = "0.00"

    # Qualified highlighting
    if "Qualified ETFs" in wb.sheetnames:

        ws = wb["Qualified ETFs"]

        for row in ws.iter_rows(
            min_row=2
        ):

            for cell in row:
                cell.fill = qualified_fill

    # Filtered highlighting
    if "Filtered ETFs" in wb.sheetnames:

        ws = wb["Filtered ETFs"]

        for row in ws.iter_rows(
            min_row=2
        ):

            for cell in row:
                cell.fill = filtered_fill

    wb.save(filename)


# ============================================================
# RUN
# ============================================================

if __name__ == "__main__":
    main()