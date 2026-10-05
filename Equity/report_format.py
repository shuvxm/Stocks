"""Formatted Excel report for backtest_asof.py - readable tables, no terminal wrap."""
import pandas as pd

SUMMARY_COLS = ["Stock", "Alert date", "Status on asof", "Setup", "Base Score",
                "Close asof", "Strategy entry", "Strategy SL", "Strategy T1"]
IND_COLS = ["Stock", "Status on asof", "Setup", "Base Score", "Close asof", "RSI",
            "ADX", "RVOL", "ATR%", "EMA20 dist %", "vs 20D high %",
            "52W high dist %", "Avg value Rs Cr/day"]
LVL_COLS = ["Stock", "Verdict", "Alert Entry", "Alert Max high after signal %",
            "Alert Result", "Alert Max gain %", "Alert Max drawdown %"]


def verdict(status):
    s = str(status)
    if s.startswith("PICK"):
        return "YES - WOULD HAVE PICKED"
    if s.startswith(("WATCHLIST", "WOULD BE WATCHLIST")):
        return "WATCHLIST"
    if s.startswith(("BORDERLINE", "WOULD BE PICK")):
        return "BORDERLINE - almost picked"
    return "NO - WOULD NOT HAVE PICKED"

def write_report(path, F, pick_rows, funnel, asof, universe, nbars):
    from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    title_font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    title_fill = PatternFill("solid", fgColor="1F4E78")
    head_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="2E75B6")
    body = Font(name="Calibri", size=10)
    wrap = Alignment(vertical="center", wrap_text=True, horizontal="left")
    center = Alignment(vertical="center", horizontal="center", wrap_text=True)
    thin = Side(style="thin", color="B4C6E7")
    bd = Border(left=thin, right=thin, top=thin, bottom=thin)
    green = PatternFill("solid", fgColor="C6EFCE")
    red = PatternFill("solid", fgColor="FFC7CE")
    amber = PatternFill("solid", fgColor="FFEB9C")

    def style(ws, title):
        nc = ws.max_column
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=nc)
        c = ws.cell(1, 1, title)
        c.font = title_font
        c.fill = title_fill
        ws.row_dimensions[1].height = 26
        ws.freeze_panes = "A3"
        ws.auto_filter.ref = f"A2:{get_column_letter(nc)}{ws.max_row}"
        for idx in range(1, nc + 1):
            letter = get_column_letter(idx)
            vals = [str(ws.cell(r, idx).value or "") for r in range(1, min(60, ws.max_row + 1))]
            ws.column_dimensions[letter].width = min(46, max(13, max(len(v) for v in vals) + 2))
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=nc):
            for cell in row:
                cell.font = body
                cell.border = bd
                cell.alignment = wrap
        for cell in ws[2]:
            cell.font = head_font
            cell.fill = head_fill
            cell.alignment = center
        ws.row_dimensions[2].height = 30

    with pd.ExcelWriter(path, engine="openpyxl") as wr:
        F = F.copy()
        F["Verdict"] = [verdict(s) for s in F["Status on asof"]]
        F[[c for c in SUMMARY_COLS if c in F]].to_excel(wr, sheet_name="SUMMARY", index=False, startrow=1)
        V = pd.DataFrame({"Stock": F["Stock"],
                          "Verdict": [verdict(s) for s in F["Status on asof"]],
                          "Status": F["Status on asof"],
                          "Setup": F["Setup"], "Score": F["Base Score"], "Why": F["Why"]})
        V.to_excel(wr, sheet_name="VERDICT", index=False, startrow=1)
        F[[c for c in IND_COLS if c in F]].to_excel(wr, sheet_name="INDICATORS", index=False, startrow=1)
        F[[c for c in LVL_COLS if c in F]].to_excel(wr, sheet_name="ALERT LEVELS", index=False, startrow=1)
        pd.DataFrame(pick_rows).to_excel(wr, sheet_name="PICKS ON DATE", index=False, startrow=1)
        pd.DataFrame(sorted(funnel.items(), key=lambda kv: -kv[1]),
                     columns=["Rejection reason", "Stocks"]).to_excel(wr, sheet_name="FUNNEL", index=False, startrow=1)
        style(wr.book["SUMMARY"], f"ALERT STOCKS as of {asof:%d-%b-%Y} | {universe} | {nbars} bars after")
        style(wr.book["VERDICT"], f"WOULD THE STRATEGY HAVE PICKED THEM? (as of {asof:%d-%b-%Y})")
        style(wr.book["INDICATORS"], f"TECHNICAL SNAPSHOT on {asof:%d-%b-%Y} (no look-ahead)")
        style(wr.book["ALERT LEVELS"], f"ALERT LEVELS outcome after {asof:%d-%b-%Y}")
        style(wr.book["PICKS ON DATE"], f"ALL PICKS/WATCH on {asof:%d-%b-%Y} (base rate)")
        style(wr.book["FUNNEL"], f"REJECTION FUNNEL on {asof:%d-%b-%Y}")
        ws = wr.book["VERDICT"]
        from openpyxl.formatting.rule import FormulaRule
        ws.conditional_formatting.add(f"B3:B{ws.max_row}",
            FormulaRule(formula=['NOT(ISERROR(SEARCH("YES",B3)))'], fill=green))
        ws.conditional_formatting.add(f"B3:B{ws.max_row}",
            FormulaRule(formula=['NOT(ISERROR(SEARCH("BORDERLINE",B3)))'], fill=amber))
        ws.conditional_formatting.add(f"B3:B{ws.max_row}",
            FormulaRule(formula=['LEFT(B3,2)="NO"'], fill=red))
        ws.column_dimensions["B"].width = 28
        ws.column_dimensions["F"].width = 46
        for sh, col in (("SUMMARY", "E"), ("INDICATORS", "D")):
            if sh in wr.book.sheetnames:
                w2 = wr.book[sh]
                w2.conditional_formatting.add(f"{col}3:{col}{w2.max_row}", ColorScaleRule(
                    start_type="num", start_value=40, start_color="F8696B",
                    mid_type="num", mid_value=60, mid_color="FFEB84",
                    end_type="num", end_value=80, end_color="63BE7B"))
        for w2 in wr.book.worksheets:
            w2.sheet_properties.pageSetUpPr.fitToPage = True

