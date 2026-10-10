"""
NEWS EFFECT REPORT - weekend news since the last NSE close, per stock
Run:  python news_effect_report.py
      python news_effect_report.py --universe NIFTY200
      python news_effect_report.py --since "2026-10-09 15:30"     (IST, overrides the automatic window)

Pipeline
  1. Window = last completed NSE session close (15:30 IST, normally Friday) -> now
  2. Pull Google News headlines for every stock in the universe, keep only those inside the window
  3. Tag each headline positive / negative by keyword, net them into a News Score per stock
  4. Add price context (last close, trend, RSI) and write NewsEffectStocks_<run date>.xlsx

Headline keywords are a rough guide to direction, not a forecast. Read the headlines before trading.
"""
import argparse
import math
import re
import sys
import time
import xml.etree.ElementTree as ET
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta, timezone
from email.utils import parsedate_to_datetime
from urllib.parse import quote_plus

import numpy as np
import pandas as pd
import requests
from openpyxl.styles import Alignment, Font, PatternFill

import indian_stock_swing_analyst as M

# ----------------------------------------------------------------------------
# CONFIG
# ----------------------------------------------------------------------------
OUTPUT_DIR = M.BASE_DIR / "output" / "news"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

IST = timezone(timedelta(hours=5, minutes=30))
UNIVERSE = "NIFTY500"
WORKERS = 8                      # parallel Google News requests
MAX_HEADLINES = 6                # headlines shown per stock in the summary sheets
UP_SCORE, DOWN_SCORE = 2, -2     # News Score needed for LIKELY UP / LIKELY DOWN

# Regex fragments (matched at a word start). The analyst's own lists are added below.
POS_STRONG = ["order win", "wins order", "bags order", "secures order", "bags", "upgrade", "buyback",
              "record profit", "profit surges", "profit jumps", "beats estimates", "usfda nod",
              "usfda approval", "fda approval", "target raised", "raises target"]
NEG_STRONG = ["fraud", "raid", "probe", "investigation", "insolvency", "default", "downgrade", "scam",
              "sebi notice", "sebi order", "warning letter", "form 483", "show cause", "net loss",
              "loss widens", "auditor"]
POS_EXTRA = ["order worth", "wins\\b", "surges?\\b", "soars?\\b", "rall(?:y|ies)\\b", "jumps?\\b", "record high",
             "52-week high", "all-time high", "profit up", "revenue up", "top pick", "buy call",
             "stocks? to buy", "tie-up", "raises stake", "nod\\b", "launch"]
NEG_EXTRA = ["falls?\\b", "drops?\\b", "slips?\\b", "crash", "tax notice", "gst notice", "tax demand",
             "penalty", "fined\\b", "strike\\b", "fire\\b", "shutdown", "stake sale", "sells stake",
             "offloads?\\b", "52-week low", "quits\\b", "resignation", "weak\\b", "declines?\\b", "cbi\\b"]

# Names the press uses that are neither the NSE symbol nor the start of the listed company name
# (ALL-CAPS entries are matched case-sensitively, see name_keys)
ALIASES = {"SBIN": ["SBI"], "LT": ["L&T"], "M&M": ["M&M", "Mahindra"], "HINDUNILVR": ["HUL"],
           "BHARTIARTL": ["Airtel"], "MARUTI": ["Maruti"], "BAJAJ-AUTO": ["Bajaj Auto"],
           "ETERNAL": ["Zomato"], "INDIGO": ["IndiGo"], "PAYTM": ["Paytm"], "NYKAA": ["Nykaa"],
           "ULTRACEMCO": ["UltraTech"], "KOTAKBANK": ["Kotak"]}
# Added to the search for one-word company names (Trent, ACC, Nava) to keep out unrelated stories
STOCK_TERMS = " (share OR shares OR stock OR NSE OR BSE)"
SAME_STORY = 0.4                 # word overlap above which two headlines count as one story
MARKET_QUERIES = ["Nifty Sensex", "stocks to watch Monday", "FII DII market"]


def _rx(words):
    return re.compile(r"\b(?:" + "|".join(words) + ")", re.I)


def _own(words):
    """The analyst's plain keywords as regex; a trailing space there means 'whole word' ("ban ")."""
    return [re.escape(w.strip()) + ("\\b" if w.endswith(" ") else "") for w in words]


RX_POS_STRONG, RX_NEG_STRONG = _rx(POS_STRONG), _rx(NEG_STRONG)
RX_POS, RX_NEG = _rx(POS_EXTRA + _own(M.POS_WORDS)), _rx(NEG_EXTRA + _own(M.NEG_WORDS))


# ----------------------------------------------------------------------------
# WINDOW
# ----------------------------------------------------------------------------
def last_close(now):
    """15:30 IST of the last completed NSE session (uses NIFTY bars, so holidays are handled)."""
    try:
        d = M.history_single(M.NIFTY)
        for ts in reversed(d.index):
            c = datetime.combine(pd.Timestamp(ts).date(), datetime.min.time(), IST).replace(hour=15, minute=30)
            if c <= now:
                return c
    except Exception:
        pass
    c = now.replace(hour=15, minute=30, second=0, microsecond=0) - timedelta(days=(now.weekday() - 4) % 7)
    return c if c <= now else c - timedelta(days=7)


# ----------------------------------------------------------------------------
# NEWS
# ----------------------------------------------------------------------------
def short_name(company):
    n = re.sub(r"\s*\(.*?\)", "", str(company or ""))
    return re.sub(r"[\s,]*\b(ltd|limited)\.?\s*$", "", n, flags=re.I).strip()


def name_keys(stock, company):
    """Strings a headline must contain to count as being about this stock. ALL-CAPS keys are case-sensitive."""
    words = short_name(company).split()
    n = 3 if len(words) > 2 and words[1].lower() in ("of", "and", "&") else 2
    keys = [" ".join(words[:n])] if words else []
    keys += [stock] if len(stock) >= 3 else []
    return [k for k in keys + ALIASES.get(stock, []) if k]


def key_regex(keys):
    # not part of a longer word, and not an exchange tag such as "BSE:XYZ" or "NSE/BSE"
    return re.compile("|".join(r"(?<![\w/])" + (re.escape(k) if k.isupper() else f"(?i:{re.escape(k)})") + r"(?![\w:/])"
                               for k in keys))


def story_words(title):
    return {w for w in re.findall(r"[a-z0-9]+", title.lower()) if len(w) > 3}


def google_news(query, since, days):
    """Returns [(datetime IST, title, source, link)] inside the window, or None if the request failed."""
    url = f"https://news.google.com/rss/search?q={quote_plus(f'{query} when:{days}d')}&hl=en-IN&gl=IN&ceid=IN:en"
    for attempt in range(3):
        try:
            r = requests.get(url, headers={"User-Agent": "Mozilla/5.0"}, timeout=15)
            if r.status_code in (429, 503):
                time.sleep(4 * (attempt + 1))
                continue
            r.raise_for_status()
            out = []
            for it in ET.fromstring(r.content).iter("item"):
                try:
                    dt = parsedate_to_datetime(it.findtext("pubDate")).astimezone(IST)
                except Exception:
                    continue
                src = (it.findtext("source") or "").strip()
                title = (it.findtext("title") or "").strip()
                if src and title.endswith(" - " + src):
                    title = title[:-len(src) - 3].strip()
                if title and dt >= since:
                    out.append((dt, title, src, (it.findtext("link") or "").strip()))
            return out
        except Exception:
            time.sleep(1 + attempt)
    return None


def tag(title):
    """(label, weight) for one headline."""
    pos = 2 if RX_POS_STRONG.search(title) else 1 if RX_POS.search(title) else 0
    neg = 2 if RX_NEG_STRONG.search(title) else 1 if RX_NEG.search(title) else 0
    if pos and neg:
        return "Mixed", 0
    return ("Positive", pos) if pos else ("Negative", -neg) if neg else ("Neutral", 0)


def stock_news(stock, company, since, days):
    """Relevant headlines for one stock, one per story, newest first. None if the fetch failed."""
    name = short_name(company) or stock
    items = google_news(f'"{name}"' + (STOCK_TERMS if " " not in name else ""), since, days)
    if items is None:
        return None
    rx = key_regex(name_keys(stock, company))
    stories, out = [], []
    for dt, title, src, link in sorted(items, reverse=True):
        w = story_words(title)
        if not rx.search(title) or any(len(w & o) >= SAME_STORY * len(w | o) for o in stories):
            continue
        stories.append(w)
        out.append((dt, title, src, link) + tag(title))
    return out


# ----------------------------------------------------------------------------
# ANALYSIS
# ----------------------------------------------------------------------------
def price_context(d):
    out = {"Last Close": np.nan, "Last Day %": np.nan, "5D %": np.nan, "RSI14": np.nan, "Trend": "N/A"}
    if d is None or len(d) < 60:
        return out
    r = M.indicators(d).iloc[-1]
    out.update({"Last Close": float(r.Close), "Last Day %": float(r.RET1), "5D %": float(r.RET5),
                "RSI14": float(r.RSI14)})
    out["Trend"] = ("Uptrend" if r.Close > r.EMA20 > r.EMA50 else
                    "Downtrend" if r.Close < r.EMA20 < r.EMA50 else "Sideways")
    return out


def expected_move(score, pos, neg):
    if score >= UP_SCORE: return "LIKELY UP"
    if score <= DOWN_SCORE: return "LIKELY DOWN"
    if score > 0: return "MILD UP"
    if score < 0: return "MILD DOWN"
    return "MIXED" if pos and neg else "NEUTRAL"


def view(score, trend):
    if score > 0:
        return {"Uptrend": "News and trend agree - upside favoured",
                "Downtrend": "Good news but trend is weak - bounce may fade"}.get(trend, "Good news, trend unclear")
    if score < 0:
        return {"Downtrend": "News and trend agree - downside favoured",
                "Uptrend": "Bad news in an uptrend - dip risk at open"}.get(trend, "Bad news, trend unclear")
    return "No clear news direction"


# ----------------------------------------------------------------------------
# EXCEL
# ----------------------------------------------------------------------------
SUM_COLS = ["Rank", "Stock", "Company", "Sector", "Expected Move", "News Score", "Positive", "Negative",
            "Total News", "Last Close", "Last Day %", "5D %", "RSI14", "Trend", "View", "Headlines"]
HEAD_COLS = ["Stock", "Published (IST)", "Sentiment", "Headline", "Source", "Link"]
MKT_COLS = ["Published (IST)", "Sentiment", "Headline", "Source", "Link"]
NEWS_WIDTH = {"Expected Move": 15, "News Score": 8, "Positive": 9, "Negative": 9, "Total News": 8, "Trend": 11,
              "View": 40, "Headlines": 110, "Headline": 100, "Published (IST)": 17, "Sentiment": 11,
              "Source": 22, "Link": 12}
FILLS = {"LIKELY UP": "C6EFCE", "MILD UP": "E2F0D9", "LIKELY DOWN": "FFC7CE", "MILD DOWN": "FCE4D6",
         "Positive": "C6EFCE", "Negative": "FFC7CE", "Mixed": "FFEB9C", "MIXED": "FFEB9C"}


def add_sheet(w, name, df, title, stamp, empty_msg):
    df.to_excel(w, sheet_name=name, index=False, startrow=3)
    ws = w.sheets[name]
    M.style_sheet(ws, df, title, stamp)
    if not len(df):
        ws["A5"] = empty_msg
        return
    for j, col in enumerate(df.columns, 1):
        letter = ws.cell(4, j).column_letter
        if col in NEWS_WIDTH:
            ws.column_dimensions[letter].width = NEWS_WIDTH[col]
        for i in range(5, ws.max_row + 1):
            cell = ws.cell(i, j)
            if col in ("News Score", "Positive", "Negative", "Total News"):
                cell.number_format = "0"
            elif col in ("Headlines", "Headline", "View"):
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            elif col == "Link" and cell.value:
                cell.hyperlink = cell.value
                cell.value = "Open"
                cell.font = Font(name=M.FONT, size=10, color="0563C1", underline="single")
            if col in ("Expected Move", "Sentiment") and cell.value in FILLS:
                cell.fill = PatternFill("solid", fgColor=FILLS[cell.value])


def write_excel(rows, heads, market, stamp, notes, path):
    def table(sel, reverse):
        part = sorted(sel, key=lambda r: (r["News Score"], r["Total News"] * (1 if reverse else -1)), reverse=reverse)
        for i, r in enumerate(part, 1):
            r["Rank"] = i
        return pd.DataFrame(part).reindex(columns=SUM_COLS)

    up = table([r for r in rows if r["News Score"] > 0], True)
    down = table([r for r in rows if r["News Score"] < 0], False)
    flat = table([r for r in rows if r["News Score"] == 0], True)
    with pd.ExcelWriter(path, engine="openpyxl") as w:
        add_sheet(w, "LIKELY UP", up, "NEWS EFFECT - stocks with positive news", stamp,
                  "No stock had net-positive news in this window.")
        add_sheet(w, "LIKELY DOWN", down, "NEWS EFFECT - stocks with negative news", stamp,
                  "No stock had net-negative news in this window.")
        add_sheet(w, "NEUTRAL", flat, "IN THE NEWS - no clear direction", stamp, "Nothing here.")
        add_sheet(w, "ALL HEADLINES", pd.DataFrame(heads).reindex(columns=HEAD_COLS),
                  "EVERY STOCK HEADLINE IN THE WINDOW", stamp, "No stock headlines found in this window.")
        add_sheet(w, "MARKET NEWS", pd.DataFrame(market).reindex(columns=MKT_COLS),
                  "MARKET-WIDE HEADLINES", stamp, "No market headlines found in this window.")
        ws = w.book.create_sheet("HOW TO READ")
        ws.column_dimensions["A"].width = 150
        for i, (txt, bold) in enumerate(notes, 1):
            c = ws.cell(i, 1, txt)
            c.font = Font(name=M.FONT, bold=bold, size=12 if bold else 10)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    return up, down


# ----------------------------------------------------------------------------
# MAIN
# ----------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--universe", choices=list(M.NSE_URLS), default=UNIVERSE)
    ap.add_argument("--since", default=None, help='window start in IST, "YYYY-MM-DD HH:MM" (default: last NSE close)')
    args = ap.parse_args()

    now = datetime.now(IST)
    since = (datetime.strptime(args.since, "%Y-%m-%d %H:%M").replace(tzinfo=IST) if args.since else last_close(now))
    days = max(1, math.ceil((now - since).total_seconds() / 86400)) + 1
    print(f"NEWS EFFECT | Universe: {args.universe} | window: {since:%a %d-%b %H:%M} -> {now:%a %d-%b %H:%M} IST")

    tickers, sectors, names = M.get_universe(args.universe)
    stocks = [M.sym(t) for t in tickers]
    done = [0]

    def work(s):
        res = stock_news(s, names.get(s, s), since, days)
        done[0] += 1
        print(f"  news {done[0]}/{len(stocks)}", end="\r")
        return s, res

    with ThreadPoolExecutor(WORKERS) as ex:
        news = dict(ex.map(work, stocks))
    failed = [s for s, v in news.items() if v is None]
    hit = {s: v for s, v in news.items() if v}
    print(f"\nHeadlines fetched for {len(stocks) - len(failed)}/{len(stocks)} stocks | {len(hit)} stocks in the news")
    if len(failed) > len(stocks) / 2:
        sys.exit("News source blocked or unreachable for most stocks - no report written.")

    data = M.download_all([s + ".NS" for s in hit]) if hit else {}
    rows, heads = [], []
    for s, items in hit.items():
        score = max(-6, min(6, sum(i[5] for i in items)))
        pos, neg = sum(i[5] > 0 for i in items), sum(i[5] < 0 for i in items)
        row = {"Stock": s, "Company": names.get(s, s), "Sector": sectors.get(s, ""),
               "Expected Move": expected_move(score, pos, neg), "News Score": score, "Positive": pos,
               "Negative": neg, "Total News": len(items)}
        row.update(price_context(data.get(s + ".NS")))
        row["View"] = view(score, row["Trend"])
        row["Headlines"] = "\n".join(f"[{i[0]:%a %H:%M}] ({i[4]}) {i[1]}" for i in items[:MAX_HEADLINES])
        rows.append(row)
        heads += [{"Stock": s, "Published (IST)": f"{i[0]:%d-%b %H:%M}", "Sentiment": i[4], "Headline": i[1],
                   "Source": i[2], "Link": i[3]} for i in items]

    market, seen = [], set()
    for q in MARKET_QUERIES:
        for dt, title, src, link in google_news(q, since, days) or []:
            if title not in seen:
                seen.add(title)
                market.append((dt, title, src, link))
    market = [{"Published (IST)": f"{dt:%d-%b %H:%M}", "Sentiment": tag(title)[0], "Headline": title,
               "Source": src, "Link": link} for dt, title, src, link in sorted(market, reverse=True)[:60]]

    m = M.market_regime()
    stamp = (f"Universe: {args.universe} | News window: {since:%a %d-%b %H:%M} -> {now:%a %d-%b %H:%M} IST | "
             f"NIFTY regime: {m['regime']} | Stocks in the news: {len(hit)}")
    notes = [
        ("HOW TO READ THIS REPORT", True),
        ("It lists every stock that had news between the last NSE close and the time of this run, and which way that news leans.", False),
        ("Several outlets repeating one story are counted once. Total News = separate stories.", False),
        ("Each headline is tagged Positive / Negative by keywords (orders, upgrades, results, buybacks vs downgrades, probes, losses, notices). "
         "Strong words count 2, ordinary ones 1. News Score = positives minus negatives, limited to -6..+6.", False),
        (f"Expected Move: LIKELY UP = News Score {UP_SCORE}+ | MILD UP = 1 | LIKELY DOWN = {DOWN_SCORE} or lower | MILD DOWN = -1 | "
         "MIXED = good and bad news cancel out | NEUTRAL = in the news, no direction.", False),
        ("Trend uses the last daily close: Uptrend = Close > EMA20 > EMA50, Downtrend = the reverse. "
         "View combines news direction with trend - the move is more reliable when both agree.", False),
        ("Keyword tagging cannot read context ('profit falls less than feared' is tagged Negative). Open the headline before acting. "
         "Good news is often priced in at the open - do not chase a big gap.", False),
        ("Not financial advice.", False),
        ("", False), ("THIS RUN", True),
        (f"Stocks checked: {len(stocks)} | news fetch failed for: {len(failed)}"
         + (f" ({', '.join(failed[:25])}{'...' if len(failed) > 25 else ''})" if failed else ""), False),
        (f"Stocks with at least one relevant headline: {len(hit)} | headlines kept: {len(heads)}", False)]

    path = OUTPUT_DIR / f"NewsEffectStocks_{now:%Y-%m-%d}.xlsx"
    up, down = write_excel(rows, heads, market, stamp, notes, path)

    def brief(df):
        return "\n".join(f"  {r['Stock']:<12} {r['Expected Move']:<12} score {int(r['News Score']):+d}  "
                         f"({r['Trend']})" for _, r in df.head(10).iterrows()) or "  none"

    summary = (f"News window: {since:%a %d-%b %H:%M} to {now:%a %d-%b %H:%M} IST\n"
               f"Stocks checked: {len(stocks)} | in the news: {len(hit)} | NIFTY regime: {m['regime']}\n\n"
               f"LIKELY UP ({len(up)} stocks with positive news)\n{brief(up)}\n\n"
               f"LIKELY DOWN ({len(down)} stocks with negative news)\n{brief(down)}\n\n"
               "Full list and every headline are in the attached Excel. Keyword-based, read the headlines before trading.\n")
    (OUTPUT_DIR / "summary.txt").write_text(summary, encoding="utf-8")
    print("\n" + summary)
    print("Excel saved:", path)


if __name__ == "__main__":
    main()
